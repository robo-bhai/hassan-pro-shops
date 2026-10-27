"""store_exports.py - small helpers for Excel (openpyxl) and PDF (reportlab) export for store pages."""
import csv
import io
from datetime import datetime

from django.http import HttpResponse


def _safe(v):
    """Protect against Excel / CSV formula injection + turn None into empty."""
    if v is None:
        return ''
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        return "'" + v
    return v


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def csv_response(filename, headers, rows, totals=None):
    resp = HttpResponse(content_type='text/csv; charset=utf-8')
    resp['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    resp.write('\ufeff')
    w = csv.writer(resp)
    w.writerow(headers)
    for r in rows:
        w.writerow([_safe(c) for c in r])
    if totals:
        w.writerow(totals)
    return resp


def xlsx_response(filename, title, headers, rows, totals=None, subtitle='', number_cols=()):
    """number_cols: column indexes (0-based) jo number format me hon."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:31]
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([subtitle or f"Generated {datetime.now():%d %b %Y %I:%M %p}"])
    ws.append([])
    ws.append(list(headers))
    hdr_row = 4
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hdr_row, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append([_safe(_num(c)) if i in number_cols else _safe(c) for i, c in enumerate(r)])
    if totals:
        ws.append([_safe(_num(c)) if i in number_cols else _safe(c) for i, c in enumerate(totals)])
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    for i in range(1, len(headers) + 1):
        width = max([len(str(ws.cell(row=r, column=i).value or '')) for r in range(hdr_row, ws.max_row + 1)] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 45)
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return resp


def pdf_response(filename, title, headers, rows, totals=None, subtitle='', landscape_mode=None, right_cols=()):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    use_landscape = landscape_mode if landscape_mode is not None else len(headers) > 6
    size = landscape(A4) if use_landscape else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle('cell', parent=styles['BodyText'], fontSize=8, leading=10)
    story = [Paragraph(f"<b>{title}</b>", styles['Title'])]
    story.append(Paragraph(subtitle or f"Generated {datetime.now():%d %b %Y %I:%M %p}", styles['Normal']))
    story.append(Spacer(1, 6 * mm))

    def esc(v):
        return str('' if v is None else v).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    data = [[Paragraph(f"<b>{esc(h)}</b>", cell) for h in headers]]
    for r in rows:
        data.append([Paragraph(esc(c), cell) for c in r])
    if totals:
        data.append([Paragraph(f"<b>{esc(c)}</b>", cell) for c in totals])
    t = Table(data, repeatRows=1)
    style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]
    t.setStyle(TableStyle(style))
    story.append(t)
    doc.build(story)
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
    return resp


def export_response(fmt, filename, title, headers, rows, totals=None, subtitle='', number_cols=()):
    """fmt: 'xlsx' | 'pdf' | 'csv'. Returns None for anything else (the view shows the normal page)."""
    if fmt == 'xlsx':
        return xlsx_response(filename, title, headers, rows, totals, subtitle, number_cols)
    if fmt == 'pdf':
        return pdf_response(filename, title, headers, rows, totals, subtitle)
    if fmt == 'csv':
        return csv_response(filename, headers, rows, totals)
    return None
