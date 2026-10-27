"""Frontend views: shareholder/cash_deductions (split from the former shareholder_cash.py)."""
from ...security.permissions import perm_required
from ...models import *
import json
import openpyxl
from reportlab.lib.units import inch
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from decimal import Decimal
from django.utils.timezone import now
from datetime import date
from django.core.paginator import Paginator
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from io import BytesIO
from django.db import transaction
from ...models import Shareholder, ShareholderCashTransaction, CompanyInfo, Purchase, SystemSetting
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from openpyxl.utils import get_column_letter
import logging
logger = logging.getLogger(__name__)


@login_required
@perm_required('manage_shareholder_cash')
def process_purchase_deduction(request, pk):
    """
    Process shareholder deduction for a specific purchase
    """
    purchase = get_object_or_404(Purchase, pk=pk)
    
    # Check permission
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied! Only administrators can process deductions.')
        return redirect('purchase_detail', pk=purchase.pk)
    
    # Check if already processed
    if purchase.shareholder_deduction_done:
        messages.warning(request, '⚠️ Deduction already processed for this purchase!')
        return redirect('purchase_deduction_details', pk=purchase.pk)
    
    # Get deduction type
    deduction_type = request.GET.get('type', 'proportional')
    
    if request.method == 'POST':
        try:
            with transaction.atomic():
                success, result = purchase.process_shareholder_deduction(request.user)
                
                if success:
                    messages.success(
                        request, 
                        f'✅ Deduction processed successfully!\n'
                        f'Total: Rs. {purchase.total_amount():,.2f}\n'
                        f'Deducted from {len(result.get("deducted_from", []))} shareholders'
                    )
                else:
                    messages.error(request, f'❌ Deduction failed: {result}')
                
                return redirect('purchase_deduction_details', pk=purchase.pk)
                
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('purchase_detail', pk=purchase.pk)
    
    # GET request - show confirmation
    shareholders = Shareholder.objects.filter(status='active')
    total_shareholders = shareholders.count()
    
    # Calculate per shareholder based on type
    total_amount = purchase.total_amount()
    if deduction_type == 'equal':
        per_shareholder = total_amount / total_shareholders if total_shareholders > 0 else 0
    else:  # proportional
        total_balance = 0
        balances = []
        for sh in shareholders:
            balance = sh.get_balance()
            balances.append(balance)
            total_balance += balance
        # Show preview
        preview_data = []
        for i, sh in enumerate(shareholders):
            if deduction_type == 'equal':
                amount = per_shareholder
            else:
                proportion = balances[i] / total_balance if total_balance > 0 else 0
                amount = total_amount * proportion
            preview_data.append({
                'name': sh.name,
                'balance': balances[i],
                'deduction': amount,
                'percentage': (amount / total_amount * 100) if total_amount > 0 else 0
            })
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'purchase': purchase,
        'total_shareholders': total_shareholders,
        'total_amount': total_amount,
        'per_shareholder': per_shareholder if deduction_type == 'equal' else None,
        'deduction_type': deduction_type,
        'preview_data': preview_data if deduction_type == 'proportional' else None,
        'is_enabled': SystemSetting.get_bool('enable_shareholder_purchase_deduction', True),
    }
    return render(request, 'shareholders/deduction/process.html', context)


@login_required
@perm_required('view_shareholders')
def purchase_deduction_details(request, pk):
    """
    View deduction details for a purchase
    """
    purchase = get_object_or_404(Purchase, pk=pk)
    
    if not purchase.shareholder_deduction_done:
        messages.warning(request, '⚠️ No deduction processed for this purchase yet!')
        return redirect('process_purchase_deduction', pk=purchase.pk)
    
    summary = purchase.get_deduction_summary()
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'purchase': purchase,
        'summary': summary,
        'deducted_from': summary.get('deducted_from', []) if summary else [],
        'failed': summary.get('failed', []) if summary else [],
        'skipped': summary.get('skipped', []) if summary else [],
        'total_shareholders': summary.get('total_shareholders', 0) if summary else 0,
        'total_amount': summary.get('total_amount', 0) if summary else 0,
        'deduction_type': purchase.shareholder_deduction_type,
        'deduction_date': purchase.shareholder_deduction_date,
    }
    return render(request, 'shareholders/deduction/details.html', context)


@login_required
@perm_required('manage_shareholder_cash')
def toggle_shareholder_deduction(request):
    """
    Toggle the shareholder deduction feature on/off
    """
    if not request.user.is_superuser:
        messages.error(request, '❌ Access denied! Only superuser can toggle this feature.')
        return redirect('dashboard')
    
    if request.method == 'POST':
        new_value = SystemSetting.toggle('enable_shareholder_purchase_deduction', request.user)
        status = 'enabled' if new_value else 'disabled'
        messages.success(request, f'✅ Shareholder purchase deduction {status}!')
        return redirect('system_settings')
    
    return redirect('system_settings')


@login_required
@perm_required('manage_shareholder_cash')
def change_deduction_type(request):
    """
    Change deduction type between equal and proportional
    """
    if not request.user.is_superuser:
        return JsonResponse({'success': False, 'message': 'Access denied!'})
    
    if request.method == 'POST':
        try:
            deduction_type = request.POST.get('deduction_type', 'proportional')
            
            if deduction_type not in ['equal', 'proportional']:
                return JsonResponse({'success': False, 'message': 'Invalid deduction type!'})
            
            SystemSetting.set_value('shareholder_deduction_type', deduction_type, request.user)
            
            return JsonResponse({
                'success': True,
                'message': f'✅ Deduction type changed to: {deduction_type}',
                'deduction_type': deduction_type
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    
    return JsonResponse({'success': False, 'message': 'Invalid request'})


@login_required
@perm_required('view_shareholders')
def shareholder_deduction_report(request):
    """
    Complete report of all shareholder deductions
    """
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    # Get all purchases with deductions
    purchases = Purchase.objects.filter(
        shareholder_deduction_done=True
    ).order_by('-shareholder_deduction_date')
    
    # Filters
    from_date = request.GET.get('from_date', '')
    to_date = request.GET.get('to_date', '')
    shareholder_id = request.GET.get('shareholder', '')
    deduction_type = request.GET.get('type', '')
    
    if from_date:
        purchases = purchases.filter(shareholder_deduction_date__date__gte=from_date)
    if to_date:
        purchases = purchases.filter(shareholder_deduction_date__date__lte=to_date)
    if deduction_type:
        purchases = purchases.filter(shareholder_deduction_type=deduction_type)
    
    # Calculate totals
    total_deductions = purchases.count()
    total_amount = Decimal('0.00')
    total_shareholders_affected = 0
    
    for p in purchases:
        data = p.shareholder_deduction_data
        if data:
            total_amount += Decimal(str(data.get('total_amount', 0)))
            total_shareholders_affected += len(data.get('deducted_from', []))
    
    # Shareholder-wise summary
    shareholder_summary = {}
    for purchase in purchases:
        data = purchase.shareholder_deduction_data
        if data and data.get('deducted_from'):
            for item in data['deducted_from']:
                name = item.get('name', 'Unknown')
                code = item.get('code', '')
                if name not in shareholder_summary:
                    shareholder_summary[name] = {
                        'name': name,
                        'code': code,
                        'total_deductions': 0,
                        'total_amount': Decimal('0.00'),
                        'purchases': []
                    }
                shareholder_summary[name]['total_deductions'] += 1
                shareholder_summary[name]['total_amount'] += Decimal(str(item.get('deducted', 0)))
                shareholder_summary[name]['purchases'].append({
                    'bill_no': purchase.bill_no,
                    'amount': Decimal(str(item.get('deducted', 0))),
                    'date': purchase.shareholder_deduction_date
                })
    
    # Sort by total amount
    shareholder_summary = sorted(
        shareholder_summary.values(),
        key=lambda x: x['total_amount'],
        reverse=True
    )
    
    # Apply shareholder filter
    if shareholder_id:
        try:
            shareholder = Shareholder.objects.get(id=shareholder_id)
            shareholder_summary = [s for s in shareholder_summary if s['name'] == shareholder.name]
        except Shareholder.DoesNotExist:
            pass
    
    # Pagination
    paginator = Paginator(purchases, 25)
    page = request.GET.get('page', 1)
    page_obj = paginator.get_page(page)
    
    # System settings
    is_enabled = SystemSetting.get_bool('enable_shareholder_purchase_deduction', True)
    current_type = SystemSetting.get_value('shareholder_deduction_type', 'proportional')
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'page_obj': page_obj,
        'purchases': page_obj,
        'total_deductions': total_deductions,
        'total_amount': total_amount,
        'total_shareholders_affected': total_shareholders_affected,
        'shareholder_summary': shareholder_summary,
        'is_enabled': is_enabled,
        'current_type': current_type,
        'shareholders': Shareholder.objects.filter(status='active'),
        'from_date': from_date,
        'to_date': to_date,
        'shareholder_id': shareholder_id,
        'deduction_type': deduction_type,
        'deduction_types': [
            ('equal', 'Equal Division'),
            ('proportional', 'Proportional by Balance'),
        ],
        'today': date.today(),
    }
    return render(request, 'shareholders/deduction/report.html', context)


@login_required
@perm_required('view_shareholders')
def shareholder_deduction_report_pdf(request):
    """
    Generate PDF report for shareholder deductions
    """
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from io import BytesIO
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    # Get data
    purchases = Purchase.objects.filter(
        shareholder_deduction_done=True
    ).order_by('-shareholder_deduction_date')[:100]
    
    total_amount = Decimal('0.00')
    for p in purchases:
        data = p.shareholder_deduction_data
        if data:
            total_amount += Decimal(str(data.get('total_amount', 0)))
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    story = []
    
    company = CompanyInfo.objects.first()
    company_name = company.name if company else "ERP SYSTEM"
    
    # Title
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, alignment=TA_CENTER)
    story.append(Paragraph(company_name, title_style))
    story.append(Paragraph("Shareholder Deduction Report", title_style))
    story.append(Paragraph(f"Generated: {now().strftime('%d-%m-%Y %I:%M %p')}", styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    # Summary
    summary_data = [
        ['Total Deductions', str(purchases.count())],
        ['Total Amount', f"Rs. {total_amount:,.2f}"],
        ['Deduction Type', SystemSetting.get_value('shareholder_deduction_type', 'proportional')],
    ]
    
    summary_table = Table(summary_data, colWidths=[2*inch, 2*inch])
    summary_table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 0.3*inch))
    
    # Deductions Table
    data = [['Bill No', 'Vendor', 'Total', 'Shareholders', 'Date']]
    for purchase in purchases:
        data.append([
            purchase.bill_no or 'N/A',
            purchase.vendor.name[:20],
            f"Rs. {purchase.total_amount():,.2f}",
            str(len(purchase.shareholder_deduction_data.get('deducted_from', [])) if purchase.shareholder_deduction_data else 0),
            purchase.shareholder_deduction_date.strftime('%d-%m-%Y') if purchase.shareholder_deduction_date else '-'
        ])
    
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1d2e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f6fa')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (3, 0), (4, -1), 'CENTER'),
    ]))
    
    story.append(table)
    story.append(Spacer(1, 0.3*inch))
    story.append(Paragraph("This is a computer generated report.", styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Shareholder_Deduction_Report_{date.today()}.pdf"'
    return response


@login_required
@perm_required('view_shareholders')
def shareholder_deduction_report_export(request):
    """
    Export shareholder deduction report to Excel
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    purchases = Purchase.objects.filter(
        shareholder_deduction_done=True
    ).order_by('-shareholder_deduction_date')[:500]
    
    wb = openpyxl.Workbook()
    
    # ========================================== #
    # SHEET 1: SUMMARY                          #
    # ========================================== #
    ws_summary = wb.active
    ws_summary.title = "Summary"
    
    company = CompanyInfo.objects.first()
    company_name = company.name if company else "ERP SYSTEM"
    
    ws_summary.merge_cells('A1:C1')
    ws_summary['A1'] = company_name
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary['A1'].alignment = Alignment(horizontal='center')
    
    ws_summary.merge_cells('A2:C2')
    ws_summary['A2'] = "Shareholder Deduction Report"
    ws_summary['A2'].font = Font(bold=True, size=14)
    ws_summary['A2'].alignment = Alignment(horizontal='center')
    
    ws_summary.merge_cells('A3:C3')
    ws_summary['A3'] = f"Generated: {now().strftime('%d-%m-%Y %I:%M %p')}"
    ws_summary['A3'].alignment = Alignment(horizontal='center')
    
    # Summary data
    total_amount = Decimal('0.00')
    total_shareholders = set()
    
    for p in purchases:
        data = p.shareholder_deduction_data
        if data:
            total_amount += Decimal(str(data.get('total_amount', 0)))
            for item in data.get('deducted_from', []):
                total_shareholders.add(item.get('name', ''))
    
    summary_data = [
        ['Metric', 'Value'],
        ['Total Deductions', len(purchases)],
        ['Total Amount', float(total_amount)],
        ['Total Shareholders Affected', len(total_shareholders)],
        ['Deduction Type', SystemSetting.get_value('shareholder_deduction_type', 'proportional')],
    ]
    
    row = 5
    for data in summary_data:
        ws_summary.cell(row=row, column=1, value=data[0])
        ws_summary.cell(row=row, column=2, value=data[1])
        if row == 5:
            ws_summary.cell(row=row, column=1).font = Font(bold=True)
            ws_summary.cell(row=row, column=2).font = Font(bold=True)
        row += 1
    
    for r in range(6, 9):
        cell = ws_summary.cell(row=r, column=2)
        if isinstance(cell.value, (int, float)):
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right')
    
    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 20
    
    # ========================================== #
    # SHEET 2: DETAILED DEDUCTIONS              #
    # ========================================== #
    ws_details = wb.create_sheet("Deductions")
    
    headers = ['#', 'Bill No', 'Vendor', 'Total Amount', 'Deduction Date', 'Shareholders']
    for col, header in enumerate(headers, 1):
        cell = ws_details.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1a1d2e", end_color="1a1d2e", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for idx, purchase in enumerate(purchases, 1):
        row = idx + 1
        ws_details.cell(row=row, column=1, value=idx)
        ws_details.cell(row=row, column=2, value=purchase.bill_no or 'N/A')
        ws_details.cell(row=row, column=3, value=purchase.vendor.name[:30])
        ws_details.cell(row=row, column=4, value=float(purchase.total_amount()))
        ws_details.cell(row=row, column=5, value=purchase.shareholder_deduction_date.strftime('%d-%m-%Y') if purchase.shareholder_deduction_date else '-')
        ws_details.cell(row=row, column=6, value=len(purchase.shareholder_deduction_data.get('deducted_from', [])) if purchase.shareholder_deduction_data else 0)
        
        ws_details.cell(row=row, column=4).number_format = '#,##0.00'
        ws_details.cell(row=row, column=4).alignment = Alignment(horizontal='right')
        ws_details.cell(row=row, column=6).alignment = Alignment(horizontal='center')
    
    # Column widths
    widths = [5, 15, 30, 15, 15, 15]
    for i, width in enumerate(widths, 1):
        ws_details.column_dimensions[get_column_letter(i)].width = width
    
    # ========================================== #
    # SHEET 3: SHAREHOLDER SUMMARY              #
    # ========================================== #
    ws_shareholder = wb.create_sheet("Shareholder Summary")
    
    # Calculate shareholder totals
    shareholder_totals = {}
    for purchase in purchases:
        data = purchase.shareholder_deduction_data
        if data and data.get('deducted_from'):
            for item in data['deducted_from']:
                name = item.get('name', 'Unknown')
                if name not in shareholder_totals:
                    shareholder_totals[name] = {
                        'code': item.get('code', ''),
                        'total_deductions': 0,
                        'total_amount': Decimal('0.00')
                    }
                shareholder_totals[name]['total_deductions'] += 1
                shareholder_totals[name]['total_amount'] += Decimal(str(item.get('deducted', 0)))
    
    sh_headers = ['#', 'Shareholder Name', 'Code', 'Total Deductions', 'Total Amount']
    for col, header in enumerate(sh_headers, 1):
        cell = ws_shareholder.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1a1d2e", end_color="1a1d2e", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    sorted_shareholders = sorted(
        shareholder_totals.items(),
        key=lambda x: x[1]['total_amount'],
        reverse=True
    )
    
    for idx, (name, data) in enumerate(sorted_shareholders, 1):
        row = idx + 1
        ws_shareholder.cell(row=row, column=1, value=idx)
        ws_shareholder.cell(row=row, column=2, value=name)
        ws_shareholder.cell(row=row, column=3, value=data['code'])
        ws_shareholder.cell(row=row, column=4, value=data['total_deductions'])
        ws_shareholder.cell(row=row, column=5, value=float(data['total_amount']))
        ws_shareholder.cell(row=row, column=5).number_format = '#,##0.00'
        ws_shareholder.cell(row=row, column=5).alignment = Alignment(horizontal='right')
        ws_shareholder.cell(row=row, column=4).alignment = Alignment(horizontal='center')
    
    sh_widths = [5, 30, 15, 18, 20]
    for i, width in enumerate(sh_widths, 1):
        ws_shareholder.column_dimensions[get_column_letter(i)].width = width
    
    # Response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Shareholder_Deduction_Report_{date.today()}.xlsx"'
    wb.save(response)
    return response


@login_required
@perm_required('manage_shareholder_cash')
def bulk_process_deductions(request):
    """
    Bulk process deductions for multiple purchases
    """
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    # ========================================== #
    # GET REQUEST - Show pending purchases      #
    # ========================================== #
    if request.method == 'GET':
        pending_purchases = Purchase.objects.filter(
            shareholder_deduction_done=False
        ).order_by('-pur_date')[:100]  # ✅ Increased limit to 100
        
        total_pending = Purchase.objects.filter(
            shareholder_deduction_done=False
        ).count()
        
        context = {
            'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
            'pending_purchases': pending_purchases,
            'total_pending': total_pending,
            'is_enabled': SystemSetting.get_bool('enable_shareholder_purchase_deduction', True),
            'deduction_type': SystemSetting.get_value('shareholder_deduction_type', 'proportional'),
        }
        return render(request, 'shareholders/deduction/bulk_process.html', context)
    
    # ========================================== #
    # POST REQUEST - Process selected purchases #
    # ========================================== #
    if request.method == 'POST':
        try:
            # Get data from request
            if request.headers.get('Content-Type') == 'application/json':
                data = json.loads(request.body)
                purchase_ids = data.get('purchase_ids', [])
            else:
                purchase_ids = request.POST.getlist('purchase_ids[]')
            
            if not purchase_ids:
                return JsonResponse({
                    'success': False, 
                    'message': 'No purchases selected!'
                })
            
            processed = 0
            failed = 0
            skipped = 0
            total_amount = Decimal('0.00')
            errors = []
            results = []
            
            with transaction.atomic():
                for purchase_id in purchase_ids:
                    try:
                        purchase = Purchase.objects.get(pk=purchase_id)
                        
                        # ✅ Skip if already processed
                        if purchase.shareholder_deduction_done:
                            skipped += 1
                            continue
                        
                        # ✅ Process deduction
                        success, result = purchase.process_shareholder_deduction(request.user)
                        
                        if success:
                            processed += 1
                            total_amount += purchase.total_amount()
                            results.append({
                                'bill_no': purchase.bill_no,
                                'amount': float(purchase.total_amount()),
                                'status': 'success'
                            })
                            logger.info(f"✅ Bulk processed: {purchase.bill_no}")
                        else:
                            failed += 1
                            errors.append(f"Purchase #{purchase_id}: {result}")
                            results.append({
                                'bill_no': purchase.bill_no,
                                'amount': float(purchase.total_amount()),
                                'status': 'failed',
                                'error': result
                            })
                            
                    except Purchase.DoesNotExist:
                        failed += 1
                        errors.append(f"Purchase #{purchase_id}: Not found")
                    except Exception as e:
                        failed += 1
                        errors.append(f"Purchase #{purchase_id}: {str(e)}")
            
            # ✅ Return JSON response for AJAX
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.headers.get('Content-Type') == 'application/json':
                return JsonResponse({
                    'success': True,
                    'message': f'✅ {processed} purchases processed! Total: Rs. {total_amount:,.2f}',
                    'processed': processed,
                    'failed': failed,
                    'skipped': skipped,
                    'total_amount': float(total_amount),
                    'errors': errors[:5],
                    'results': results
                })
            
            # ✅ Return HTML response for form submission
            messages.success(
                request, 
                f'✅ {processed} purchases processed! Total: Rs. {total_amount:,.2f}. Failed: {failed}'
            )
            return redirect('shareholder_deduction_report')
            
        except Exception as e:
            logger.error(f"Bulk processing error: {e}")
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False, 
                    'message': str(e)
                })
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('bulk_process_deductions')
    
    return redirect('bulk_process_deductions')


@login_required
@perm_required('view_shareholders')
def shareholder_deduction_history(request, pk):
    """
    View deduction history for a specific shareholder
    """
    shareholder = get_object_or_404(Shareholder, pk=pk)
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    # Find all purchases where this shareholder was deducted
    deduction_history = []
    total_deducted = Decimal('0.00')
    
    purchases = Purchase.objects.filter(
        shareholder_deduction_done=True
    ).order_by('-shareholder_deduction_date')
    
    for purchase in purchases:
        data = purchase.shareholder_deduction_data
        if data and data.get('deducted_from'):
            for item in data['deducted_from']:
                if item.get('name') == shareholder.name:
                    deduction_history.append({
                        'purchase': purchase,
                        'amount': Decimal(str(item.get('deducted', 0))),
                        'percentage': item.get('percentage', 0),
                        'balance_before': Decimal(str(item.get('balance_before', 0))),
                        'balance_after': Decimal(str(item.get('balance_after', 0))),
                        'date': purchase.shareholder_deduction_date,
                    })
                    total_deducted += Decimal(str(item.get('deducted', 0)))
                    break
    
    # Get all transactions for this shareholder
    transactions = ShareholderCashTransaction.objects.filter(
        shareholder=shareholder
    ).order_by('-created_at')[:50]
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'shareholder': shareholder,
        'deduction_history': deduction_history,
        'total_deducted': total_deducted,
        'total_deductions': len(deduction_history),
        'transactions': transactions,
        'current_balance': shareholder.get_balance(),
        'total_shares': shareholder.total_shares(),
    }
    return render(request, 'shareholders/deduction_history.html', context)
