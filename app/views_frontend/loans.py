"""Frontend views: loans."""
from ..models import *
import openpyxl
from reportlab.lib.units import inch
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Q
from decimal import Decimal
from django.utils.timezone import now
from datetime import date, timedelta, datetime
from django.core.paginator import Paginator
from django.contrib import messages
from django.http import HttpResponse
from io import BytesIO
from ..models import CashBalance, CompanyInfo, Expense, LoanGiven, LoanGivenPayment
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from openpyxl.utils import get_column_letter


@login_required
def loan_pay(request, pk):
    """Make loan payment"""
    
    loan = get_object_or_404(ShareholderLoan, pk=pk)
    
    # Check permission
    if not request.user.is_superuser and not request.user.is_staff:
        if not hasattr(request.user, 'shareholder_profile'):
            messages.error(request, 'Access denied!')
            return redirect('dashboard')
        if loan.shareholder != request.user.shareholder_profile:
            messages.error(request, 'Access denied!')
            return redirect('dashboard')
    
    if request.method == 'POST':
        try:
            amount = Decimal(request.POST.get('amount', 0))
            
            if amount <= 0:
                messages.error(request, 'Amount must be greater than zero!')
                return redirect('loan_pay', pk=pk)
            
            # Check if amount exceeds total due
            total_due = loan.outstanding + loan.interest_accrued
            if amount > total_due:
                messages.error(request, f'Amount cannot exceed total due: Rs. {total_due:,.2f}')
                return redirect('loan_pay', pk=pk)
            
            # Accrue interest first
            loan.accrue_interest()
            
            # Make payment
            loan.make_payment(
                amount=amount,
                payment_method=request.POST.get('payment_method', 'Cash'),
                reference_no=request.POST.get('reference_no', '')
            )
            
            # Update loan status
            if loan.outstanding <= 0 and loan.interest_accrued <= 0:
                loan.status = 'paid'
                loan.save()
            
            messages.success(request, f'✅ Payment of Rs. {amount:,.2f} made successfully!')
            return redirect('loan_detail', pk=loan.pk)
            
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan': loan,
        'total_due': loan.outstanding + loan.interest_accrued,
    }
    return render(request, 'shareholders/loan/pay.html', context)


@login_required
def loan_statement(request, pk):
    """Generate loan statement PDF"""
    
    loan = get_object_or_404(ShareholderLoan, pk=pk)
    
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    story = []
    
    company = CompanyInfo.objects.first()
    company_name = company.name if company else "ERP SYSTEM"
    
    # Title
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, alignment=TA_CENTER)
    story.append(Paragraph(company_name, title_style))
    story.append(Paragraph("LOAN STATEMENT", title_style))
    story.append(Paragraph(f"Loan No: {loan.loan_no}", styles['Normal']))
    story.append(Spacer(1, 0.3*inch))
    
    # Loan Info
    info_data = [
        ['Shareholder:', loan.shareholder.name],
        ['Loan Type:', loan.get_loan_type_display()],
        ['Principal:', f"Rs. {loan.principal:,.2f}"],
        ['Interest Rate:', f"{loan.interest_rate}%"],
        ['Start Date:', loan.start_date.strftime('%d-%m-%Y')],
        ['Maturity Date:', loan.maturity_date.strftime('%d-%m-%Y')],
        ['Status:', loan.get_status_display()],
    ]
    
    info_table = Table(info_data, colWidths=[2*inch, 4*inch])
    info_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.3*inch))
    
    # Payment Summary
    interest_payments = InterestPayment.objects.filter(loan=loan)
    principal_payments = PrincipalPayment.objects.filter(loan=loan)
    
    payment_data = [
        ['Description', 'Amount (Rs.)'],
        ['Total Principal', f"{loan.principal:,.2f}"],
        ['Principal Paid', f"{loan.principal - loan.outstanding:,.2f}"],
        ['Principal Outstanding', f"{loan.outstanding:,.2f}"],
        ['', ''],
        ['Interest Accrued', f"{loan.interest_accrued:,.2f}"],
        ['Interest Paid', f"{loan.interest_paid:,.2f}"],
        ['Total Due', f"{loan.outstanding + loan.interest_accrued:,.2f}"],
    ]
    
    payment_table = Table(payment_data, colWidths=[3*inch, 2*inch])
    payment_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1d2e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (1, 1), (1, -1), 'RIGHT'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e8ecf4')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    story.append(payment_table)
    story.append(Spacer(1, 0.3*inch))
    
    # Payment History
    if interest_payments or principal_payments:
        story.append(Paragraph("<b>Payment History</b>", styles['Heading2']))
        history_data = [['Date', 'Type', 'Amount', 'Method']]
        
        for payment in interest_payments:
            history_data.append([
                payment.payment_date.strftime('%d-%m-%Y'),
                'Interest',
                f"Rs. {payment.amount:,.2f}",
                payment.payment_method or '-'
            ])
        
        for payment in principal_payments:
            history_data.append([
                payment.payment_date.strftime('%d-%m-%Y'),
                'Principal',
                f"Rs. {payment.amount:,.2f}",
                payment.payment_method or '-'
            ])
        
        history_table = Table(history_data, repeatRows=1)
        history_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1d2e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        story.append(history_table)
    
    doc.build(story)
    buffer.seek(0)
    
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Loan_Statement_{loan.loan_no}.pdf"'
    return response


@login_required
def loan_return_list(request):
    """List all loan returns"""
    loan_returns = LoanReturn.objects.all().order_by('-return_date')
    
    search = request.GET.get('search', '')
    if search:
        loan_returns = loan_returns.filter(
            Q(return_no__icontains=search) |
            Q(lender_name__icontains=search) |
            Q(loan_agreement_no__icontains=search)
        )
    
    status = request.GET.get('status', '')
    if status:
        loan_returns = loan_returns.filter(status=status)
    
    paginator = Paginator(loan_returns, 25)
    page = request.GET.get('page', 1)
    page_obj = paginator.get_page(page)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'page_obj': page_obj,
        'loan_returns': page_obj,
        'total_returns': loan_returns.count(),
        'total_amount': loan_returns.aggregate(total=Sum('total_amount'))['total'] or 0,
        'pending_count': loan_returns.filter(status='pending').count(),
        'paid_count': loan_returns.filter(status='paid').count(),
    }
    return render(request, 'loan_returns/list.html', context)


@login_required
def loan_return_create(request):
    """Create new loan return - Auto-fill lender name"""
    
    # ✅ GET loan_id from URL parameter
    loan_id = request.GET.get('loan')
    pre_filled_loan = None
    pre_filled_lender = ''
    pre_filled_agreement = ''
    pre_filled_amount = 0
    
    if loan_id:
        try:
            pre_filled_loan = Loan.objects.get(pk=loan_id)
            pre_filled_lender = pre_filled_loan.lender_name  # ✅ Auto-fill lender
            pre_filled_agreement = pre_filled_loan.agreement_no or pre_filled_loan.loan_no
            pre_filled_amount = pre_filled_loan.remaining_amount
        except Loan.DoesNotExist:
            pass
    
    if request.method == 'POST':
        try:
            # ✅ SAFE DECIMAL CONVERSION
            def safe_decimal(value, default=0):
                try:
                    if not value or value == '':
                        return Decimal(default)
                    cleaned = str(value).replace(',', '').replace(' ', '').replace('Rs.', '').replace('$', '')
                    if cleaned == '':
                        return Decimal(default)
                    return Decimal(cleaned)
                except:
                    return Decimal(default)
            
            # ✅ Get loan agreement from hidden field
            loan_agreement_no = request.POST.get('loan_agreement_no', '').strip()
            
            if not loan_agreement_no:
                loan_id_from_post = request.POST.get('loan_id')
                if loan_id_from_post:
                    try:
                        loan = Loan.objects.get(pk=loan_id_from_post)
                        loan_agreement_no = loan.agreement_no or loan.loan_no
                    except Loan.DoesNotExist:
                        pass
            
            if not loan_agreement_no:
                messages.error(request, '❌ Loan Agreement No is required!')
                return redirect('loan_return_create')
            
            lender_name = request.POST.get('lender_name', '').strip()
            if not lender_name:
                messages.error(request, '❌ Lender name is required!')
                return redirect('loan_return_create')
            
            amount = safe_decimal(request.POST.get('amount', 0))
            if amount <= 0:
                messages.error(request, '❌ Amount must be greater than zero!')
                return redirect('loan_return_create')
            
            return_date = request.POST.get('return_date')
            if not return_date:
                messages.error(request, '❌ Return date is required!')
                return redirect('loan_return_create')
            
            # ✅ Create loan return
            loan_return = LoanReturn.objects.create(
                loan_agreement_no=loan_agreement_no,
                lender_name=lender_name,
                return_type=request.POST.get('return_type', 'installment'),
                amount=amount,
                principal_amount=0,
                interest_amount=0,
                late_fee=0,
                return_date=return_date,
                due_date=request.POST.get('due_date') or None,
                payment_method=request.POST.get('payment_method', 'bank_transfer'),
                reference_no=request.POST.get('reference_no', ''),
                description=request.POST.get('description', ''),
                notes=request.POST.get('notes', ''),
                created_by=request.user,
                status='pending'
            )
            
            # Process payment immediately
            if request.POST.get('process_now') == 'on':
                try:
                    loan_return.process_payment(request.user)
                    messages.success(
                        request, 
                        f'✅ Loan Return #{loan_return.return_no} created and paid!\n'
                        f'Amount: Rs. {loan_return.total_amount:,.2f} deducted from main cash.'
                    )
                except ValidationError as e:
                    messages.error(request, f'❌ Payment failed: {str(e)}')
                    return redirect('loan_return_detail', pk=loan_return.pk)
            else:
                messages.success(
                    request, 
                    f'✅ Loan Return #{loan_return.return_no} created!\n'
                    f'Amount: Rs. {loan_return.total_amount:,.2f} (Pending)'
                )
            
            return redirect('loan_return_detail', pk=loan_return.pk)
            
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
            print(f"LOAN RETURN ERROR: {e}")
            import traceback
            traceback.print_exc()
            return redirect('loan_return_create')
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'return_types': LoanReturn.RETURN_TYPES,
        'payment_methods': Expense.PAYMENT_METHODS,
        'today': date.today().strftime('%Y-%m-%d'),
        'current_balance': CashBalance.get_balance(),
        'pre_filled_loan': pre_filled_loan,
        'pre_filled_lender': pre_filled_lender,  # ✅ Pass to template
        'pre_filled_agreement': pre_filled_agreement,
        'pre_filled_amount': pre_filled_amount,
        'loan_id': loan_id,
    }
    return render(request, 'loan_returns/create.html', context)


@login_required
def loan_return_detail(request, pk):
    """View loan return detail"""
    loan_return = get_object_or_404(LoanReturn, pk=pk)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan_return': loan_return,
        'can_pay': loan_return.status == 'pending',
        'can_delete': loan_return.status == 'pending',
    }
    return render(request, 'loan_returns/detail.html', context)


@login_required
def loan_return_pay(request, pk):
    """Process loan return payment"""
    loan_return = get_object_or_404(LoanReturn, pk=pk)
    
    if loan_return.status == 'paid':
        messages.warning(request, '⚠️ This loan return is already paid!')
        return redirect('loan_return_detail', pk=pk)
    
    if request.method == 'POST':
        try:
            loan_return.process_payment(request.user)
            messages.success(
                request, 
                f'✅ Loan Return #{loan_return.return_no} paid!\n'
                f'Amount: Rs. {loan_return.total_amount:,.2f} deducted from main cash.'
            )
            return redirect('loan_return_detail', pk=pk)
        except ValidationError as e:
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('loan_return_detail', pk=pk)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan_return': loan_return,
        'current_balance': CashBalance.get_balance(),
    }
    return render(request, 'loan_returns/pay.html', context)


@login_required
def loan_return_delete(request, pk):
    """Delete loan return"""
    loan_return = get_object_or_404(LoanReturn, pk=pk)
    
    if loan_return.status == 'paid':
        messages.error(request, '❌ Cannot delete a paid loan return!')
        return redirect('loan_return_detail', pk=pk)
    
    if request.method == 'POST':
        return_no = loan_return.return_no
        loan_return.delete()
        messages.success(request, f'🗑️ Loan Return #{return_no} deleted!')
        return redirect('loan_return_list')
    
    return redirect('loan_return_detail', pk=pk)


@login_required
def loan_list(request):
    """List all loans"""
    from django.db.models import Sum, Q
    from django.core.paginator import Paginator
    
    loans = Loan.objects.all().order_by('-loan_date')
    
    # Update loan balances
    for loan in loans:
        total_paid = LoanReturn.objects.filter(
            loan_agreement_no=loan.agreement_no or loan.loan_no,
            status='paid'
        ).aggregate(total=Sum('total_amount'))['total'] or 0
        
        loan.total_payable = loan.principal_amount + (loan.principal_amount * loan.interest_rate / 100)
        loan.paid_amount = total_paid
        loan.remaining_amount = loan.total_payable - total_paid
        
        paid_count = LoanReturn.objects.filter(
            loan_agreement_no=loan.agreement_no or loan.loan_no,
            status='paid'
        ).count()
        loan.paid_installments = paid_count
        
        if loan.remaining_amount <= 0 and loan.total_payable > 0:
            loan.status = 'paid'
        
        loan.save(update_fields=['paid_amount', 'remaining_amount', 'paid_installments', 'status'])
    
    # Filters
    search = request.GET.get('search', '')
    if search:
        loans = loans.filter(
            Q(loan_no__icontains=search) |
            Q(lender_name__icontains=search) |
            Q(agreement_no__icontains=search)
        )
    
    status = request.GET.get('status', '')
    if status:
        loans = loans.filter(status=status)
    
    loan_type = request.GET.get('type', '')
    if loan_type:
        loans = loans.filter(loan_type=loan_type)
    
    # Statistics
    total_loans = loans.count()
    total_principal = loans.aggregate(total=Sum('principal_amount'))['total'] or 0
    total_paid = loans.aggregate(total=Sum('paid_amount'))['total'] or 0
    total_remaining = loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    active_loans = loans.filter(status='active').count()
    
    # Pagination
    paginator = Paginator(loans, 25)
    page = request.GET.get('page', 1)
    page_obj = paginator.get_page(page)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'page_obj': page_obj,
        'loans': page_obj,
        'search': search,
        'status': status,
        'loan_type': loan_type,
        'total_loans': total_loans,
        'total_principal': total_principal,
        'total_paid': total_paid,
        'total_remaining': total_remaining,
        'active_loans': active_loans,
        'loan_types': Loan.LOAN_TYPES,
        'status_choices': Loan.LOAN_STATUS,
    }
    
    # ✅ IMPORTANT: Return must be here
    return render(request, 'loans/list.html', context)


@login_required
def loan_create(request):
    """Create new loan - Simplified (No Agreement No, No Interest Rate)"""
    
    if request.method == 'POST':
        try:
            # ✅ SAFE DECIMAL CONVERSION
            def safe_decimal(value, default=0):
                try:
                    if not value or value == '':
                        return Decimal(default)
                    cleaned = str(value).replace(',', '').replace(' ', '').replace('Rs.', '').replace('$', '')
                    if cleaned == '':
                        return Decimal(default)
                    return Decimal(cleaned)
                except:
                    return Decimal(default)
            
            # Validation
            lender_name = request.POST.get('lender_name', '').strip()
            if not lender_name:
                messages.error(request, '❌ Lender/Bank name is required!')
                return redirect('loan_create')
            
            principal = safe_decimal(request.POST.get('principal_amount', 0))
            if principal <= 0:
                messages.error(request, '❌ Principal amount must be greater than zero!')
                return redirect('loan_create')
            
            loan_date = request.POST.get('loan_date')
            if not loan_date:
                messages.error(request, '❌ Loan date is required!')
                return redirect('loan_create')
            
            due_date = request.POST.get('due_date')
            if not due_date:
                messages.error(request, '❌ Due date is required!')
                return redirect('loan_create')
            
            # ✅ SAFE INTEGER CONVERSION - FIXED
            try:
                total_installments = int(request.POST.get('total_installments', 0) or 0)
            except ValueError:
                total_installments = 0
            
            # Create loan
            loan = Loan.objects.create(
                loan_type=request.POST.get('loan_type', 'business'),
                lender_name=lender_name,
                lender_contact=request.POST.get('lender_contact', ''),
                principal_amount=principal,
                interest_rate=0,
                loan_date=loan_date,
                due_date=due_date,
                total_installments=total_installments,
                purpose=request.POST.get('purpose', ''),
                notes=request.POST.get('notes', ''),
                created_by=request.user,
                status='active',
                paid_amount=0,
                remaining_amount=principal,
                total_payable=principal,
                paid_installments=0,
            )
            
            # Process loan - ADD to main cash
            if request.POST.get('process_now') == 'on':
                try:
                    loan.process_loan_received(request.user)
                    messages.success(
                        request, 
                        f'✅ Loan #{loan.loan_no} created and cash added!\n'
                        f'Amount: Rs. {loan.principal_amount:,.2f} added to main cash.'
                    )
                except Exception as e:
                    messages.error(request, f'❌ Cash addition failed: {str(e)}')
                    return redirect('loan_detail', pk=loan.pk)
            else:
                messages.success(
                    request, 
                    f'✅ Loan #{loan.loan_no} created!\n'
                    f'Amount: Rs. {loan.principal_amount:,.2f} (Pending)'
                )
            
            # Save agreement file (if uploaded)
            if request.FILES.get('agreement_file'):
                loan.agreement_file = request.FILES['agreement_file']
                loan.save()
            
            return redirect('loan_detail', pk=loan.pk)
            
        except ValidationError as e:
            messages.error(request, f'❌ Validation Error: {str(e)}')
            return redirect('loan_create')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
            print(f"LOAN CREATE ERROR: {e}")
            import traceback
            traceback.print_exc()
            return redirect('loan_create')
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan_types': Loan.LOAN_TYPES,
        'today': date.today().strftime('%Y-%m-%d'),
        'current_balance': CashBalance.get_balance(),
    }
    return render(request, 'loans/create.html', context)


@login_required
def loan_detail(request, pk):
    """View loan details"""
    loan = get_object_or_404(Loan, pk=pk)
    
    # ✅ FIX: Calculate loan returns
    from django.db.models import Sum
    
    # Get loan returns for this loan
    loan_returns = LoanReturn.objects.filter(
        loan_agreement_no=loan.agreement_no or loan.loan_no
    ).order_by('-return_date')
    
    total_returned = loan_returns.filter(status='paid').aggregate(
        total=Sum('total_amount')
    )['total'] or 0
    
    # ✅ FIX: Update loan totals
    loan.total_payable = loan.principal_amount + (loan.principal_amount * loan.interest_rate / 100)
    loan.paid_amount = total_returned
    loan.remaining_amount = loan.total_payable - total_returned
    
    # ✅ FIX: Count paid installments
    paid_installments = loan_returns.filter(status='paid').count()
    loan.paid_installments = paid_installments
    
    # ✅ FIX: Update status
    if loan.remaining_amount <= 0 and loan.total_payable > 0:
        loan.status = 'paid'
        loan.save(update_fields=['status', 'paid_amount', 'remaining_amount', 'paid_installments'])
    
    # ✅ Calculate progress
    progress_percent = 0
    if loan.total_payable > 0:
        progress_percent = (total_returned / loan.total_payable) * 100
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan': loan,
        'loan_returns': loan_returns,
        'total_returned': total_returned,
        'remaining_balance': loan.remaining_amount,
        'can_add_return': loan.status == 'active',
        'can_delete': loan.status == 'active',
        'progress_percent': progress_percent,
        'paid_installments': paid_installments,
    }
    return render(request, 'loans/detail.html', context)


@login_required
def loan_delete(request, pk):
    """Delete loan"""
    loan = get_object_or_404(Loan, pk=pk)
    
    if loan.status == 'paid':
        messages.error(request, '❌ Cannot delete a paid loan!')
        return redirect('loan_detail', pk=pk)
    
    if request.method == 'POST':
        loan_no = loan.loan_no
        loan.delete()
        messages.success(request, f'🗑️ Loan #{loan_no} deleted!')
        return redirect('loan_list')
    
    return redirect('loan_detail', pk=pk)


@login_required
def loan_dashboard(request):
    """Loan Management Dashboard"""
    from django.db.models import Sum, Count
    from django.utils.timezone import now
    from datetime import timedelta
    
    loans = Loan.objects.all()
    
    # ✅ Statistics
    total_loans = loans.count()
    active_loans = loans.filter(status='active').count()
    paid_loans = loans.filter(status='paid').count()
    
    total_principal = loans.aggregate(total=Sum('principal_amount'))['total'] or 0
    total_paid = loans.aggregate(total=Sum('paid_amount'))['total'] or 0
    total_remaining = loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    
    # ✅ Paid percentage
    paid_percent = 0
    if total_principal > 0:
        paid_percent = (total_paid / total_principal) * 100
    
    # ✅ Overdue loans
    today = now().date()
    overdue_loans = loans.filter(
        status='active',
        due_date__lt=today
    )
    overdue_count = overdue_loans.count()
    overdue_amount = overdue_loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    
    # ✅ Loan type distribution
    loan_type_data = []
    loan_type_labels = []
    for loan_type, label in Loan.LOAN_TYPES:
        count = loans.filter(loan_type=loan_type).count()
        if count > 0:
            loan_type_labels.append(label)
            loan_type_data.append(count)
    
    # ✅ Monthly payments (last 6 months)
    monthly_labels = []
    monthly_payments = []
    for i in range(6):
        month = today.replace(day=1) - timedelta(days=30 * i)
        month_name = month.strftime('%b %Y')
        monthly_labels.insert(0, month_name)
        
        total = LoanReturn.objects.filter(
            status='paid',
            return_date__year=month.year,
            return_date__month=month.month
        ).aggregate(total=Sum('total_amount'))['total'] or 0
        monthly_payments.insert(0, float(total))
    
    # ✅ Recent activities
    recent_activities = []
    
    # Recent loan creations
    for loan in loans.order_by('-created_at')[:3]:
        recent_activities.append({
            'type': 'loan_created',
            'message': f'New loan #{loan.loan_no} created with {loan.lender_name} (Rs. {loan.principal_amount:,.2f})',
            'time_ago': f'{loan.created_at.strftime("%d-%m-%Y")}'
        })
    
    # Recent payments
    for payment in LoanReturn.objects.filter(status='paid').order_by('-updated_at')[:3]:
        recent_activities.append({
            'type': 'payment_made',
            'message': f'Payment of Rs. {payment.total_amount:,.2f} made to {payment.lender_name}',
            'time_ago': f'{payment.updated_at.strftime("%d-%m-%Y")}'
        })
    
    # Overdue alerts
    for loan in overdue_loans[:3]:
        days_overdue = (today - loan.due_date).days
        recent_activities.append({
            'type': 'overdue',
            'message': f'Loan #{loan.loan_no} from {loan.lender_name} is {days_overdue} days overdue',
            'time_ago': f'{days_overdue} days'
        })
    
    # ✅ Upcoming payments
    upcoming_payments = []
    for loan in loans.filter(status='active').order_by('due_date'):
        if loan.due_date >= today:
            days_left = (loan.due_date - today).days
            status = 'due_soon' if days_left <= 7 else 'upcoming'
            status_display = '🔴 Due Soon' if days_left <= 7 else '🟢 Upcoming'
            if days_left <= 30:  # Only show next 30 days
                upcoming_payments.append({
                    'lender': loan.lender_name,
                    'agreement_no': loan.agreement_no or loan.loan_no,
                    'amount': loan.remaining_amount,
                    'status': status,
                    'status_display': status_display,
                    'days_left': days_left
                })
                if len(upcoming_payments) >= 5:
                    break
    
    # ✅ Lender wise summary
    lender_summary = []
    for lender in loans.values('lender_name').distinct():
        lender_loans = loans.filter(lender_name=lender['lender_name'])
        lender_summary.append({
            'name': lender['lender_name'],
            'count': lender_loans.count(),
            'total_principal': lender_loans.aggregate(total=Sum('principal_amount'))['total'] or 0,
            'total_paid': lender_loans.aggregate(total=Sum('paid_amount'))['total'] or 0,
            'total_remaining': lender_loans.aggregate(total=Sum('remaining_amount'))['total'] or 0,
        })
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'total_loans': total_loans,
        'active_loans': active_loans,
        'paid_loans': paid_loans,
        'total_principal': total_principal,
        'total_paid': total_paid,
        'total_remaining': total_remaining,
        'paid_percent': paid_percent,
        'overdue_count': overdue_count,
        'overdue_amount': overdue_amount,
        'loan_type_labels': loan_type_labels,
        'loan_type_data': loan_type_data,
        'monthly_labels': monthly_labels,
        'monthly_payments': monthly_payments,
        'recent_activities': recent_activities[:10],
        'upcoming_payments': upcoming_payments,
        'lender_summary': lender_summary,
    }
    return render(request, 'loans/dashboard.html', context)


@login_required
def loan_export_report(request):
    """Export Loan Report to Excel"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    # Get all loans
    loans = Loan.objects.all().order_by('-loan_date')
    
    # Create workbook
    wb = openpyxl.Workbook()
    
    # ========================================== #
    # SHEET 1: Loan Summary                     #
    # ========================================== #
    ws1 = wb.active
    ws1.title = "Loan Summary"
    
    # Title
    company = CompanyInfo.objects.first()
    company_name = company.name if company else "ERP SYSTEM"
    
    ws1.merge_cells('A1:H1')
    ws1['A1'] = company_name
    ws1['A1'].font = Font(bold=True, size=16)
    ws1['A1'].alignment = Alignment(horizontal='center')
    
    ws1.merge_cells('A2:H2')
    ws1['A2'] = "LOAN MANAGEMENT REPORT"
    ws1['A2'].font = Font(bold=True, size=14)
    ws1['A2'].alignment = Alignment(horizontal='center')
    
    ws1.merge_cells('A3:H3')
    ws1['A3'] = f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws1['A3'].alignment = Alignment(horizontal='center')
    
    # Headers
    headers = ['#', 'Loan No', 'Lender', 'Type', 'Principal', 'Paid', 'Remaining', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=5, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1a1d2e", end_color="1a1d2e", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Data
    row = 6
    total_principal = 0
    total_paid = 0
    total_remaining = 0
    
    for idx, loan in enumerate(loans, 1):
        ws1.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws1.cell(row=row, column=2, value=loan.loan_no)
        ws1.cell(row=row, column=3, value=loan.lender_name)
        ws1.cell(row=row, column=4, value=loan.get_loan_type_display())
        ws1.cell(row=row, column=5, value=float(loan.principal_amount))
        ws1.cell(row=row, column=5).number_format = '#,##0.00'
        ws1.cell(row=row, column=6, value=float(loan.paid_amount))
        ws1.cell(row=row, column=6).number_format = '#,##0.00'
        ws1.cell(row=row, column=7, value=float(loan.remaining_amount))
        ws1.cell(row=row, column=7).number_format = '#,##0.00'
        ws1.cell(row=row, column=8, value=loan.get_status_display())
        ws1.cell(row=row, column=8).alignment = Alignment(horizontal='center')
        
        total_principal += loan.principal_amount
        total_paid += loan.paid_amount
        total_remaining += loan.remaining_amount
        row += 1
    
    # Totals row
    ws1.cell(row=row, column=4, value="TOTAL").font = Font(bold=True)
    ws1.cell(row=row, column=5, value=float(total_principal)).font = Font(bold=True)
    ws1.cell(row=row, column=5).number_format = '#,##0.00'
    ws1.cell(row=row, column=6, value=float(total_paid)).font = Font(bold=True)
    ws1.cell(row=row, column=6).number_format = '#,##0.00'
    ws1.cell(row=row, column=7, value=float(total_remaining)).font = Font(bold=True)
    ws1.cell(row=row, column=7).number_format = '#,##0.00'
    
    # ========================================== #
    # SHEET 2: Loan Returns                      #
    # ========================================== #
    ws2 = wb.create_sheet("Loan Returns")
    
    headers2 = ['#', 'Return No', 'Lender', 'Agreement No', 'Amount', 'Date', 'Type', 'Status']
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1a1d2e", end_color="1a1d2e", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    returns = LoanReturn.objects.all().order_by('-return_date')
    row = 2
    total_returned = 0
    
    for idx, return_obj in enumerate(returns, 1):
        ws2.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws2.cell(row=row, column=2, value=return_obj.return_no)
        ws2.cell(row=row, column=3, value=return_obj.lender_name)
        ws2.cell(row=row, column=4, value=return_obj.loan_agreement_no)
        ws2.cell(row=row, column=5, value=float(return_obj.total_amount))
        ws2.cell(row=row, column=5).number_format = '#,##0.00'
        ws2.cell(row=row, column=6, value=return_obj.return_date.strftime('%d-%m-%Y'))
        ws2.cell(row=row, column=7, value=return_obj.get_return_type_display())
        ws2.cell(row=row, column=8, value=return_obj.get_status_display())
        ws2.cell(row=row, column=8).alignment = Alignment(horizontal='center')
        
        total_returned += return_obj.total_amount
        row += 1
    
    ws2.cell(row=row, column=4, value="TOTAL").font = Font(bold=True)
    ws2.cell(row=row, column=5, value=float(total_returned)).font = Font(bold=True)
    ws2.cell(row=row, column=5).number_format = '#,##0.00'
    
    # ========================================== #
    # SHEET 3: Lender Summary                    #
    # ========================================== #
    ws3 = wb.create_sheet("Lender Summary")
    
    headers3 = ['#', 'Lender', 'Total Loans', 'Principal', 'Paid', 'Remaining', 'Status']
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1a1d2e", end_color="1a1d2e", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    lender_summary = []
    for lender in loans.values('lender_name').distinct():
        lender_loans = loans.filter(lender_name=lender['lender_name'])
        lender_summary.append({
            'name': lender['lender_name'],
            'count': lender_loans.count(),
            'principal': lender_loans.aggregate(total=Sum('principal_amount'))['total'] or 0,
            'paid': lender_loans.aggregate(total=Sum('paid_amount'))['total'] or 0,
            'remaining': lender_loans.aggregate(total=Sum('remaining_amount'))['total'] or 0,
        })
    
    row = 2
    for idx, lender in enumerate(lender_summary, 1):
        ws3.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws3.cell(row=row, column=2, value=lender['name'])
        ws3.cell(row=row, column=3, value=lender['count']).alignment = Alignment(horizontal='center')
        ws3.cell(row=row, column=4, value=float(lender['principal']))
        ws3.cell(row=row, column=4).number_format = '#,##0.00'
        ws3.cell(row=row, column=5, value=float(lender['paid']))
        ws3.cell(row=row, column=5).number_format = '#,##0.00'
        ws3.cell(row=row, column=6, value=float(lender['remaining']))
        ws3.cell(row=row, column=6).number_format = '#,##0.00'
        ws3.cell(row=row, column=7, value="✅ Paid" if lender['remaining'] <= 0 else "⏳ Active")
        ws3.cell(row=row, column=7).alignment = Alignment(horizontal='center')
        row += 1
    
    # Auto-fit columns for all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[col_letter].width = adjusted_width
    
    # Response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"Loan_Report_{date.today().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def loan_export_pdf(request):
    """Export Loan Report to PDF"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loans = Loan.objects.all().order_by('-loan_date')
    returns = LoanReturn.objects.all().order_by('-return_date')
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=landscape(A4),
        topMargin=0.3*inch,
        bottomMargin=0.3*inch,
        leftMargin=0.3*inch,
        rightMargin=0.3*inch
    )
    
    styles = getSampleStyleSheet()
    story = []
    
    company = CompanyInfo.objects.first()
    company_name = company.name if company else "ERP SYSTEM"
    
    # Title
    title_style = ParagraphStyle(
        'Title',
        parent=styles['Heading1'],
        fontSize=16,
        alignment=TA_CENTER,
        spaceAfter=5
    )
    
    story.append(Paragraph(company_name, title_style))
    story.append(Paragraph("LOAN MANAGEMENT REPORT", title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}", styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    # Summary Stats
    total_loans = loans.count()
    total_principal = loans.aggregate(total=Sum('principal_amount'))['total'] or 0
    total_paid = loans.aggregate(total=Sum('paid_amount'))['total'] or 0
    total_remaining = loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    active_loans = loans.filter(status='active').count()
    paid_loans = loans.filter(status='paid').count()
    
    summary_data = [
        ['Total Loans', 'Active', 'Paid', 'Total Principal', 'Total Paid', 'Total Remaining'],
        [
            str(total_loans),
            str(active_loans),
            str(paid_loans),
            f"Rs. {total_principal:,.2f}",
            f"Rs. {total_paid:,.2f}",
            f"Rs. {total_remaining:,.2f}"
        ]
    ]
    
    summary_table = Table(summary_data, colWidths=[1.2*inch]*6)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1d2e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, 1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    
    story.append(summary_table)
    story.append(Spacer(1, 0.2*inch))
    
    # Loan List Table
    story.append(Paragraph("<b>Loan List</b>", styles['Heading2']))
    
    loan_data = [['#', 'Loan No', 'Lender', 'Type', 'Principal', 'Paid', 'Remaining', 'Status']]
    
    for idx, loan in enumerate(loans, 1):
        loan_data.append([
            str(idx),
            loan.loan_no,
            loan.lender_name[:20],
            loan.get_loan_type_display()[:15],
            f"Rs. {loan.principal_amount:,.2f}",
            f"Rs. {loan.paid_amount:,.2f}",
            f"Rs. {loan.remaining_amount:,.2f}",
            loan.get_status_display()
        ])
    
    loan_table = Table(loan_data, repeatRows=1)
    loan_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1d2e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f6fa')]),
        ('ALIGN', (4, 0), (6, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    
    story.append(loan_table)
    story.append(Spacer(1, 0.2*inch))
    
    # Footer
    story.append(Paragraph("This is a computer generated report.", styles['Normal']))
    story.append(Paragraph(f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}", styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    
    response = HttpResponse(buffer, content_type='application/pdf')
    filename = f"Loan_Report_{date.today().strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def loan_given_dashboard(request):
    """Loan Given Dashboard"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loans = LoanGiven.objects.all()
    
    # Statistics
    total_loans = loans.count()
    active_loans = loans.filter(status='active').count()
    paid_loans = loans.filter(status='paid').count()
    defaulted_loans = loans.filter(status='defaulted').count()
    
    total_principal = loans.aggregate(total=Sum('principal_amount'))['total'] or 0
    total_received = loans.aggregate(total=Sum('received_amount'))['total'] or 0
    total_remaining = loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    
    received_percent = 0
    if total_principal > 0:
        received_percent = (total_received / total_principal) * 100
    
    # Overdue loans
    today = date.today()
    overdue_loans = loans.filter(status='active', due_date__lt=today)
    overdue_count = overdue_loans.count()
    overdue_amount = overdue_loans.aggregate(total=Sum('remaining_amount'))['total'] or 0
    
    # Recent payments
    recent_payments = LoanGivenPayment.objects.select_related('loan_given').order_by('-created_at')[:10]
    
    # Upcoming payments
    upcoming_payments = []
    for loan in loans.filter(status='active').order_by('due_date'):
        if loan.due_date >= today:
            days_left = (loan.due_date - today).days
            if days_left <= 30:
                upcoming_payments.append({
                    'borrower': loan.borrower_name,
                    'loan_no': loan.loan_no,
                    'amount': loan.remaining_amount,
                    'days_left': days_left,
                    'status': '🔴 Due Soon' if days_left <= 7 else '🟢 Upcoming'
                })
                if len(upcoming_payments) >= 10:
                    break
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'total_loans': total_loans,
        'active_loans': active_loans,
        'paid_loans': paid_loans,
        'defaulted_loans': defaulted_loans,
        'total_principal': total_principal,
        'total_received': total_received,
        'total_remaining': total_remaining,
        'received_percent': received_percent,
        'overdue_count': overdue_count,
        'overdue_amount': overdue_amount,
        'recent_payments': recent_payments,
        'upcoming_payments': upcoming_payments,
        'current_balance': CashBalance.get_balance(),
    }
    return render(request, 'loans_given/dashboard.html', context)


@login_required
def loan_given_list(request):
    """List all loans given"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loans = LoanGiven.objects.all().order_by('-loan_date')
    
    # Filters
    search = request.GET.get('search', '')
    if search:
        loans = loans.filter(
            Q(loan_no__icontains=search) |
            Q(borrower_name__icontains=search) |
            Q(borrower_contact__icontains=search)
        )
    
    status = request.GET.get('status', '')
    if status:
        loans = loans.filter(status=status)
    
    loan_type = request.GET.get('type', '')
    if loan_type:
        loans = loans.filter(loan_type=loan_type)
    
    # Statistics
    total_loans = loans.count()
    total_principal = loans.aggregate(total=Sum('principal_amount'))['total'] or 0
    total_received = loans.aggregate(total=Sum('received_amount'))['total'] or 0
    total_remaining = total_principal - total_received
    
    paginator = Paginator(loans, 25)
    page = request.GET.get('page', 1)
    page_obj = paginator.get_page(page)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'page_obj': page_obj,
        'loans': page_obj,
        'search': search,
        'status': status,
        'loan_type': loan_type,
        'total_loans': total_loans,
        'total_principal': total_principal,
        'total_received': total_received,
        'total_remaining': total_remaining,
        'loan_types': LoanGiven.LOAN_TYPES,
        'status_choices': LoanGiven.LOAN_STATUS,
    }
    return render(request, 'loans_given/list.html', context)


@login_required
def loan_given_create(request):
    """Create new loan given"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    if request.method == 'POST':
        try:
            # Safe Decimal conversion
            def safe_decimal(value, default=0):
                try:
                    if not value or value == '':
                        return Decimal(default)
                    cleaned = str(value).replace(',', '').replace(' ', '').replace('Rs.', '').replace('$', '')
                    if cleaned == '':
                        return Decimal(default)
                    return Decimal(cleaned)
                except:
                    return Decimal(default)
            
            # Validation
            borrower_name = request.POST.get('borrower_name', '').strip()
            if not borrower_name:
                messages.error(request, '❌ Borrower name is required!')
                return redirect('loan_given_create')
            
            principal = safe_decimal(request.POST.get('principal_amount', 0))
            if principal <= 0:
                messages.error(request, '❌ Principal amount must be greater than zero!')
                return redirect('loan_given_create')
            
            loan_date = request.POST.get('loan_date')
            if not loan_date:
                messages.error(request, '❌ Loan date is required!')
                return redirect('loan_given_create')
            
            due_date = request.POST.get('due_date')
            if not due_date:
                messages.error(request, '❌ Due date is required!')
                return redirect('loan_given_create')
            
            # Create loan
            loan = LoanGiven.objects.create(
                loan_type=request.POST.get('loan_type', 'personal'),
                borrower_name=borrower_name,
                borrower_contact=request.POST.get('borrower_contact', ''),
                borrower_cnic=request.POST.get('borrower_cnic', ''),
                borrower_address=request.POST.get('borrower_address', ''),
                principal_amount=principal,
                interest_rate=safe_decimal(request.POST.get('interest_rate', 0)),
                loan_date=loan_date,
                due_date=due_date,
                total_installments=int(request.POST.get('total_installments', 0) or 0),
                guarantor_name=request.POST.get('guarantor_name', ''),
                guarantor_contact=request.POST.get('guarantor_contact', ''),
                notes=request.POST.get('notes', ''),
                created_by=request.user,
                status='active'
            )
            
            # Process loan - Deduct from main cash
            if request.POST.get('process_now') == 'on':
                try:
                    loan.process_loan_given(request.user)
                    messages.success(
                        request, 
                        f'✅ Loan #{loan.loan_no} created and cash deducted!\n'
                        f'Amount: Rs. {loan.principal_amount:,.2f} deducted from main cash.'
                    )
                except Exception as e:
                    messages.error(request, f'❌ Cash deduction failed: {str(e)}')
                    return redirect('loan_given_detail', pk=loan.pk)
            else:
                messages.success(
                    request, 
                    f'✅ Loan #{loan.loan_no} created!\n'
                    f'Amount: Rs. {loan.principal_amount:,.2f} (Pending)'
                )
            
            # Save agreement file
            if request.FILES.get('agreement_file'):
                loan.agreement_file = request.FILES['agreement_file']
                loan.save()
            
            return redirect('loan_given_detail', pk=loan.pk)
            
        except ValidationError as e:
            messages.error(request, f'❌ Validation Error: {str(e)}')
            return redirect('loan_given_create')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('loan_given_create')
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan_types': LoanGiven.LOAN_TYPES,
        'today': date.today().strftime('%Y-%m-%d'),
        'current_balance': CashBalance.get_balance(),
    }
    return render(request, 'loans_given/create.html', context)


@login_required
def loan_given_detail(request, pk):
    """View loan given details"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loan = get_object_or_404(LoanGiven, pk=pk)
    payments = loan.payments.all().order_by('-payment_date')
    
    # Calculate totals
    total_received = payments.aggregate(total=Sum('amount'))['total'] or 0
    
    # Update loan totals
    loan.received_amount = total_received
    loan.remaining_amount = loan.total_receivable - total_received
    
    if loan.remaining_amount <= 0 and loan.total_receivable > 0:
        loan.status = 'paid'
        loan.paid_date = now().date()
        loan.save(update_fields=['received_amount', 'remaining_amount', 'status', 'paid_date'])
    
    # Calculate progress
    progress_percent = 0
    if loan.total_receivable > 0:
        progress_percent = (total_received / loan.total_receivable) * 100
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan': loan,
        'payments': payments,
        'total_received': total_received,
        'remaining_balance': loan.remaining_amount,
        'progress_percent': progress_percent,
        'can_receive_payment': loan.status == 'active',
        'can_delete': loan.status != 'paid',
    }
    return render(request, 'loans_given/detail.html', context)


@login_required
def loan_given_pay(request, pk):
    """Receive payment from borrower"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loan = get_object_or_404(LoanGiven, pk=pk)
    
    if loan.status == 'paid':
        messages.warning(request, '⚠️ This loan is already fully paid!')
        return redirect('loan_given_detail', pk=pk)
    
    if request.method == 'POST':
        try:
            amount = Decimal(request.POST.get('amount', 0))
            payment_method = request.POST.get('payment_method', 'cash')
            reference_no = request.POST.get('reference_no', '')
            notes = request.POST.get('notes', '')
            
            if amount <= 0:
                messages.error(request, '❌ Amount must be greater than zero!')
                return redirect('loan_given_pay', pk=pk)
            
            if amount > loan.remaining_amount:
                messages.error(
                    request, 
                    f'❌ Amount cannot exceed remaining balance: Rs. {loan.remaining_amount:,.2f}'
                )
                return redirect('loan_given_pay', pk=pk)
            
            # Process payment
            loan.receive_payment(amount, request.user, payment_method, reference_no)
            
            # Add notes
            if notes:
                payment = LoanGivenPayment.objects.filter(
                    loan_given=loan,
                    amount=amount,
                    payment_date=now().date()
                ).first()
                if payment:
                    payment.notes = notes
                    payment.save()
            
            messages.success(
                request, 
                f'✅ Payment of Rs. {amount:,.2f} received from {loan.borrower_name}!\n'
                f'Remaining: Rs. {loan.remaining_amount:,.2f}'
            )
            return redirect('loan_given_detail', pk=loan.pk)
            
        except ValidationError as e:
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('loan_given_pay', pk=pk)
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')
            return redirect('loan_given_pay', pk=pk)
    
    context = {
        'company_name': CompanyInfo.objects.first().name if CompanyInfo.objects.exists() else 'ERP System',
        'loan': loan,
        'remaining_balance': loan.remaining_amount,
        'today': date.today().strftime('%Y-%m-%d'),
    }
    return render(request, 'loans_given/pay.html', context)


@login_required
def loan_given_delete(request, pk):
    """Delete loan given"""
    
    if not request.user.is_superuser and not request.user.is_staff:
        messages.error(request, '❌ Access denied!')
        return redirect('dashboard')
    
    loan = get_object_or_404(LoanGiven, pk=pk)
    
    if loan.status == 'paid':
        messages.error(request, '❌ Cannot delete a paid loan!')
        return redirect('loan_given_detail', pk=pk)
    
    if request.method == 'POST':
        loan_no = loan.loan_no
        loan.delete()
        messages.success(request, f'🗑️ Loan #{loan_no} deleted!')
        return redirect('loan_given_list')
    
    return redirect('loan_given_detail', pk=pk)
